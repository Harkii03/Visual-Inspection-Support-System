from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from animal_assistant.workbook import (
    AnimalWorkbook,
    animal_code,
    capture_sort_key,
    device_folder,
    is_reviewed,
    normalized_manual_values,
    sequence_key_and_frame,
)
import app as app_module


def make_batch(tmp_path: Path) -> Path:
    batch = tmp_path / "003001_260721"
    image_dir = batch / "server" / "003001_260721_inf_serow"
    image_dir.mkdir(parents=True)
    (image_dir / "003_001_test.jpg").write_bytes(b"not-a-real-jpeg")

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet"
    sheet["B2"] = "装置番号: 003001"
    sheet["C2"] = "解析画像数: 1/1"
    sheet["D2"] = "解析日: 260721"
    headers = [
        None,
        "ファイル名",
        "ファイル更新日時",
        "動物名",
        "目視による動物名",
        "数",
        "目視による数",
        "デバイス",
        "最大確信度",
    ]
    for column, value in enumerate(headers, 1):
        sheet.cell(3, column).value = value
    sheet.cell(4, 2).value = "003_001_test.jpg"
    sheet.cell(4, 4).value = "serow（カモシカ）"
    sheet.cell(4, 6).value = 1
    sheet.cell(4, 8).value = "サーバー"
    sheet.cell(4, 9).value = 0.95
    sheet.cell(4, 9).number_format = "0.000"
    book.save(batch / "result_003001_260721_all.xlsx")
    return batch


class WorkbookTests(unittest.TestCase):
    def test_path_components(self):
        self.assertEqual(animal_code("serow（カモシカ）"), "serow")
        self.assertEqual(device_folder("サーバー"), "server")
        self.assertEqual(device_folder("装置"), "device")
        self.assertEqual(
            sequence_key_and_frame("003_004_2026.06.12.1939_4.jpeg"),
            ("003_004_2026.06.12.1939", 4),
        )
        self.assertEqual(capture_sort_key("003_004_2026.06.12.1939_4.jpeg")[1], 4)

    def test_review_rules(self):
        self.assertEqual(normalized_manual_values("serow（カモシカ）", "serow（カモシカ）", "1"), (None, 1))
        self.assertEqual(normalized_manual_values("serow（カモシカ）", "boar（イノシシ）", 2), ("boar（イノシシ）", 2))
        self.assertEqual(normalized_manual_values("serow（カモシカ）", "いない", ""), ("いない", None))
        self.assertTrue(is_reviewed(None, 1))
        self.assertTrue(is_reviewed("いない", None))
        self.assertFalse(is_reviewed(None, None))
        with self.assertRaises(ValueError):
            normalized_manual_values("serow（カモシカ）", "serow（カモシカ）", "1.5")

    def test_open_resolve_and_save_without_touching_other_cells(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            book = AnimalWorkbook.open_batch(batch)
            self.assertEqual(book.device_number, "003001")
            self.assertEqual(book.analysis_date, "260721")
            self.assertTrue(book.image_path_for_row(book.rows[0]).is_file())
            self.assertEqual(book.completed_count(), 0)

            book.save_review(0, "serow（カモシカ）", 1)
            self.assertIsNotNone(book.backup_path)
            self.assertTrue(book.backup_path.is_file())
            book.close()

            saved = load_workbook(batch / "result_003001_260721_all.xlsx")
            sheet = saved["Sheet"]
            self.assertIsNone(sheet.cell(4, 5).value)
            self.assertEqual(sheet.cell(4, 7).value, 1)
            self.assertEqual(sheet.cell(4, 9).value, 0.95)
            self.assertEqual(sheet.cell(4, 9).number_format, "0.000")
            saved.close()

    def test_org_image_is_used_when_inf_image_is_absent(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            inf_image = batch / "server" / "003001_260721_inf_serow" / "003_001_test.jpg"
            org_dir = batch / "server" / "003001_260721_org_serow"
            org_dir.mkdir()
            org_image = org_dir / inf_image.name
            inf_image.replace(org_image)

            book = AnimalWorkbook.open_batch(batch)
            self.assertEqual(book.image_path_for_row(book.rows[0]), org_image)
            self.assertTrue(book.row_data(0)["imageExists"])
            book.close()

    def test_inf_image_with_jpg_extension_is_preferred_for_excel_jpeg(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            workbook_path = batch / "result_003001_260721_all.xlsx"
            excel = load_workbook(workbook_path)
            excel.active.cell(4, 2).value = "003_001_test.jpeg"
            excel.save(workbook_path)
            excel.close()

            inf_dir = batch / "server" / "003001_260721_inf_serow"
            jpg_image = inf_dir / "003_001_test.jpg"
            org_dir = batch / "server" / "003001_260721_org_serow"
            org_dir.mkdir()
            (org_dir / "003_001_test.jpeg").write_bytes(b"org")

            book = AnimalWorkbook.open_batch(batch)
            self.assertEqual(book.image_path_for_row(book.rows[0]), jpg_image)
            book.close()

    def test_context_images_include_adjacent_frames_across_animal_folders(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            workbook_path = batch / "result_003001_260721_all.xlsx"
            excel = load_workbook(workbook_path)
            sheet = excel.active
            sheet.cell(4, 2).value = "003_001_2026.07.06.1309_1.jpg"
            for row, frame, animal in ((5, 0, "boar（イノシシ）"), (6, 2, "serow（カモシカ）")):
                sheet.cell(row, 2).value = f"003_001_2026.07.06.1309_{frame}.jpg"
                sheet.cell(row, 4).value = animal
                sheet.cell(row, 6).value = 1
                sheet.cell(row, 8).value = "サーバー"
            excel.save(workbook_path)
            excel.close()

            serow_dir = batch / "server" / "003001_260721_inf_serow"
            (serow_dir / "003_001_test.jpg").replace(
                serow_dir / "003_001_2026.07.06.1309_1.jpg"
            )
            (serow_dir / "003_001_2026.07.06.1309_2.jpg").write_bytes(b"frame2")
            boar_dir = batch / "server" / "003001_260721_inf_boar"
            boar_dir.mkdir()
            (boar_dir / "003_001_2026.07.06.1309_0.jpg").write_bytes(b"frame0")

            book = AnimalWorkbook.open_batch(batch)
            context = book.context_images(0)
            self.assertEqual([item["frame"] for item in context], [0, 1, 2])
            self.assertEqual(sum(item["isCurrent"] for item in context), 1)
            self.assertTrue(all(item["imageExists"] for item in context))
            book.close()

    def test_context_images_include_nearby_capture_times(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            workbook_path = batch / "result_003001_260721_all.xlsx"
            excel = load_workbook(workbook_path)
            sheet = excel.active
            for row, time_value in ((4, "1309"), (5, "1308"), (6, "1310")):
                sheet.cell(row, 2).value = f"003_001_2026.07.06.{time_value}_0.jpg"
                sheet.cell(row, 4).value = "serow（カモシカ）"
                sheet.cell(row, 6).value = 1
                sheet.cell(row, 8).value = "サーバー"
            excel.save(workbook_path)
            excel.close()

            image_dir = batch / "server" / "003001_260721_inf_serow"
            (image_dir / "003_001_test.jpg").replace(
                image_dir / "003_001_2026.07.06.1309_0.jpg"
            )
            for time_value in ("1308", "1310"):
                (image_dir / f"003_001_2026.07.06.{time_value}_0.jpg").write_bytes(b"nearby")

            book = AnimalWorkbook.open_batch(batch)
            context = book.context_images(0)
            self.assertEqual(
                [item["capturedAt"] for item in context],
                ["2026-07-06 13:08", "2026-07-06 13:09", "2026-07-06 13:10"],
            )
            book.close()

    def test_context_images_include_undetected_files_not_in_excel(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            workbook_path = batch / "result_003001_260721_all.xlsx"
            excel = load_workbook(workbook_path)
            excel.active.cell(4, 2).value = "003_001_2026.07.06.1309_0.jpg"
            excel.save(workbook_path)
            excel.close()

            inf_dir = batch / "server" / "003001_260721_inf_serow"
            (inf_dir / "003_001_test.jpg").replace(
                inf_dir / "003_001_2026.07.06.1309_0.jpg"
            )
            server_root = batch / "server"
            for time_value in ("1308", "1310"):
                (server_root / f"003_001_2026.07.06.{time_value}_0.jpg").write_bytes(b"undetected")

            book = AnimalWorkbook.open_batch(batch)
            context = book.context_images(0)
            self.assertEqual(len(context), 3)
            self.assertEqual([item["detected"] for item in context], [False, True, False])
            self.assertIn("1308", book.context_image_path(0, 0).name)
            book.close()

    def test_absent_saves_blank_count(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            book = AnimalWorkbook.open_batch(batch)
            book.save_review(0, "いない", "")
            self.assertEqual(book.completed_count(), 1)
            book.close()

            saved = load_workbook(batch / "result_003001_260721_all.xlsx")
            sheet = saved["Sheet"]
            self.assertEqual(sheet.cell(4, 5).value, "いない")
            self.assertIsNone(sheet.cell(4, 7).value)
            saved.close()

    def test_reference_sample_is_never_used_as_test_target(self):
        sample = Path(r"C:\Users\yudai\Desktop\画像解析260721\003001_260721")
        self.assertIn("Desktop", str(sample))
        self.assertFalse(Path(__file__).resolve().is_relative_to(sample))

    def test_local_api_flow_uses_generated_copy(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            client = app_module.app.test_client()

            opened = client.post("/api/open", json={"path": str(batch)})
            self.assertEqual(opened.status_code, 200)
            self.assertEqual(opened.get_json()["total"], 1)

            row = client.get("/api/row/0")
            self.assertEqual(row.status_code, 200)
            self.assertTrue(row.get_json()["row"]["imageExists"])

            image = client.get("/api/image/0")
            self.assertEqual(image.status_code, 200)
            image.close()

            saved = client.post(
                "/api/save",
                json={"index": 0, "animal": "serow（カモシカ）", "count": 1},
            )
            self.assertEqual(saved.status_code, 200)
            self.assertEqual(saved.get_json()["completed"], 1)

            if app_module.session is not None:
                app_module.session.close()
                app_module.session = None

    def test_api_explains_excel_file_lock(self):
        with __import__("tempfile").TemporaryDirectory() as temporary:
            batch = make_batch(Path(temporary))
            client = app_module.app.test_client()
            opened = client.post("/api/open", json={"path": str(batch)})
            self.assertEqual(opened.status_code, 200)

            original_save = app_module.session._save_atomic
            locked = PermissionError(5, "Access is denied")
            locked.winerror = 5
            app_module.session._save_atomic = lambda: (_ for _ in ()).throw(locked)
            try:
                response = client.post(
                    "/api/save",
                    json={"index": 0, "animal": "serow（カモシカ）", "count": 1},
                )
                self.assertEqual(response.status_code, 400)
                self.assertIn("対象Excelを閉じて", response.get_json()["error"])
            finally:
                app_module.session._save_atomic = original_save
                app_module.session.close()
                app_module.session = None


if __name__ == "__main__":
    unittest.main()
